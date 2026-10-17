"""Importa a planilha Controle_Financeiro (.xlsx) para o banco, sem duplicar dados."""

import re
from zipfile import BadZipFile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Product, Transaction, User
from app.schemas import (
    ImportIssue,
    ImportReport,
    PlanSettingsIn,
    PlanSettingsOut,
    ProductIn,
    TransactionIn,
)
from app.services import plan_settings
from app.services import spreadsheet_layout as layout


class SpreadsheetError(ValueError):
    pass


@dataclass
class _Collected:
    errors: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)

    def error(self, sheet: str, row: int, message: str) -> None:
        self.errors.append(ImportIssue(sheet=sheet, row=row, message=message))

    def warn(self, sheet: str, row: int, message: str) -> None:
        self.warnings.append(ImportIssue(sheet=sheet, row=row, message=message))


# ---------- Conversão de células ----------
def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _text(value: Any) -> str:
    if value is None or _is_formula(value):
        return ""
    return str(value).strip()


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


_MONEY_CHARS = re.compile(r"[^\d,.\-]")


def parse_money(value: Any) -> float | None:
    """Aceita números e textos como 'R$ 2,568.19' ou 'R$ 2.568,19'."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _MONEY_CHARS.sub("", _text(value))
    if not text:
        return None
    if "," in text and "." in text:
        decimal_sep = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal_sep == "," else ","
        text = text.replace(thousands, "").replace(decimal_sep, ".")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _column_list(sheet: Worksheet, column: str, first_row: int) -> list[str]:
    values = []
    row = first_row
    while (text := _text(sheet[f"{column}{row}"].value)):
        values.append(text)
        row += 1
    return values


def _row_values(sheet: Worksheet, row: int, width: int) -> list[Any]:
    return [sheet.cell(row=row, column=col).value for col in range(1, width + 1)]


def _is_blank(values: list[Any]) -> bool:
    return all(_text(v) == "" for v in values)


def _first_error(exc: ValidationError) -> str:
    detail = exc.errors()[0]
    where = ".".join(str(p) for p in detail.get("loc", ()))
    return f"{where}: {detail['msg']}" if where else detail["msg"]


# ---------- Leitura das abas ----------
def _read_settings(book, base: PlanSettingsIn, out: _Collected) -> PlanSettingsIn | None:
    if layout.SHEET_CONFIG not in book.sheetnames:
        return None
    config = book[layout.SHEET_CONFIG]
    data = base.model_dump()

    start = _as_date(config[layout.CONFIG_PLAN_START].value)
    end = _as_date(config[layout.CONFIG_PLAN_END].value)
    if start:
        data["plan_start"] = start
    if end:
        data["plan_end"] = end
    days = parse_money(config[layout.CONFIG_DAYS_SOON].value)
    if days is not None:
        data["days_soon"] = int(days)
    limit = parse_money(config[layout.CONFIG_TIGHT_LIMIT].value)
    if limit is not None:
        data["tight_limit"] = limit

    product_categories = _column_list(
        config, layout.CONFIG_PRODUCT_CATEGORIES_COL, layout.CONFIG_LIST_FIRST_ROW
    )
    transaction_categories = _column_list(
        config, layout.CONFIG_TRANSACTION_CATEGORIES_COL, layout.CONFIG_LIST_FIRST_ROW
    )
    if product_categories:
        data["product_categories"] = product_categories
    if len(transaction_categories) > layout.INCOME_CATEGORY_COUNT:
        data["income_categories"] = transaction_categories[: layout.INCOME_CATEGORY_COUNT]
        data["expense_categories"] = transaction_categories[layout.INCOME_CATEGORY_COUNT :]

    if layout.SHEET_INCOME in book.sheetnames:
        income = book[layout.SHEET_INCOME]
        for key, cell in layout.INCOME_BUDGET_CELLS.items():
            raw = income[cell].value
            value = parse_money(raw)
            if value is None and _text(raw):
                out.warn(layout.SHEET_INCOME, int(cell[1:]), f"Valor '{raw}' ignorado")
            elif value is not None:
                data[key] = value
        base_label = _text(income[layout.INCOME_PLAN_BASE].value).casefold()
        data["plan_base"] = layout.invert(layout.BASE_LABELS).get(base_label, data["plan_base"])

    if layout.SHEET_SIMULATION in book.sheetnames:
        simulation = book[layout.SHEET_SIMULATION]
        scenarios = [parse_money(simulation[c].value) for c in layout.SIMULATION_SCENARIO_CELLS]
        if all(v is not None for v in scenarios):
            data["scenarios"] = scenarios

    try:
        return PlanSettingsIn(**data)
    except ValidationError as exc:
        out.error(layout.SHEET_CONFIG, 0, _first_error(exc))
        return None


def _read_products(book, out: _Collected) -> list[ProductIn]:
    if layout.SHEET_PRODUCTS not in book.sheetnames:
        return []
    sheet = book[layout.SHEET_PRODUCTS]
    priorities = layout.invert(layout.PRIORITY_LABELS)
    items: list[ProductIn] = []

    for row in range(layout.FIRST_DATA_ROW, sheet.max_row + 1):
        name, category, price, priority, target, saved = _row_values(sheet, row, 6)
        if _is_blank([name, category, price, priority, target, saved]):
            continue
        priority_code = priorities.get(_text(priority).casefold())
        if priority_code is None:
            # Na planilha, prioridade desconhecida entra no fim da fila (= Baixa).
            out.warn(layout.SHEET_PRODUCTS, row, f"Prioridade '{_text(priority)}' tratada como Baixa")
            priority_code = "low"
        try:
            items.append(
                ProductIn(
                    name=_text(name),
                    category=_text(category),
                    price=parse_money(price),
                    priority=priority_code,
                    target_date=_as_date(target),
                    saved=parse_money(saved) or 0,
                )
            )
        except ValidationError as exc:
            out.error(layout.SHEET_PRODUCTS, row, _first_error(exc))
    return items


def _read_transactions(
    book, categories: PlanSettingsIn, out: _Collected
) -> list[TransactionIn]:
    if layout.SHEET_TRANSACTIONS not in book.sheetnames:
        return []
    sheet = book[layout.SHEET_TRANSACTIONS]
    kinds = layout.invert(layout.KIND_LABELS)
    groups = layout.invert(layout.GROUP_LABELS)
    statuses = layout.invert(layout.STATUS_LABELS)
    allowed = {
        "income": {c.casefold() for c in categories.income_categories},
        "expense": {c.casefold() for c in categories.expense_categories},
    }
    items: list[TransactionIn] = []

    for row in range(layout.FIRST_DATA_ROW, sheet.max_row + 1):
        values = _row_values(sheet, row, 7)
        if _is_blank(values):
            continue
        when, description, kind, group, category, amount, status = values

        kind_code = kinds.get(_text(kind).casefold())
        parsed_date = _as_date(when)
        if parsed_date is None:
            out.error(layout.SHEET_TRANSACTIONS, row, "Data ausente ou inválida")
            continue
        if kind_code is None:
            out.error(layout.SHEET_TRANSACTIONS, row, f"Tipo '{_text(kind)}' inválido")
            continue

        group_text = _text(group)
        group_code = groups.get(group_text.casefold()) if group_text else None
        if kind_code == "income" and group_text:
            out.warn(layout.SHEET_TRANSACTIONS, row, "Receita com Grupo: grupo ignorado")
        elif group_text and group_code is None:
            out.warn(layout.SHEET_TRANSACTIONS, row, f"Grupo '{group_text}' tratado como Variável")

        category_text = _text(category) or "Outros"
        if category_text.casefold() not in allowed[kind_code]:
            out.warn(
                layout.SHEET_TRANSACTIONS,
                row,
                f"Categoria '{category_text}' não é de {layout.KIND_LABELS[kind_code].lower()}",
            )

        try:
            items.append(
                TransactionIn(
                    date=parsed_date,
                    description=_text(description),
                    kind=kind_code,
                    expense_group=group_code,
                    category=category_text,
                    amount=parse_money(amount),
                    status=statuses.get(_text(status).casefold(), "paid"),
                )
            )
        except ValidationError as exc:
            out.error(layout.SHEET_TRANSACTIONS, row, _first_error(exc))
    return items


# ---------- Deduplicação ----------
def _transaction_key(item: TransactionIn | Transaction) -> tuple:
    return (
        item.date,
        item.description.casefold(),
        item.kind,
        round(float(item.amount), 2),
        item.category.casefold(),
    )


def _new_transactions(
    db: Session, user: User, incoming: list[TransactionIn]
) -> list[TransactionIn]:
    """Mantém repetições legítimas: só importa o que excede o que já existe no banco."""
    existing = Counter(
        _transaction_key(t)
        for t in db.scalars(select(Transaction).where(Transaction.user_id == user.id))
    )
    fresh = []
    for item in incoming:
        key = _transaction_key(item)
        if existing[key] > 0:
            existing[key] -= 1
        else:
            fresh.append(item)
    return fresh


def _new_products(db: Session, user: User, incoming: list[ProductIn]) -> list[ProductIn]:
    seen = {
        name.casefold()
        for name in db.scalars(select(Product.name).where(Product.user_id == user.id))
    }
    fresh = []
    for item in incoming:
        key = item.name.casefold()
        if key not in seen:
            seen.add(key)
            fresh.append(item)
    return fresh


# ---------- Entrada pública ----------
def import_workbook(db: Session, user: User, content: bytes, dry_run: bool) -> ImportReport:
    try:
        book = load_workbook(BytesIO(content), data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise SpreadsheetError("Arquivo inválido: envie a planilha no formato .xlsx") from exc

    known = {layout.SHEET_TRANSACTIONS, layout.SHEET_PRODUCTS, layout.SHEET_CONFIG}
    if not known.intersection(book.sheetnames):
        raise SpreadsheetError(
            "A planilha não tem as abas Lançamentos, Produtos ou Config"
        )

    out = _Collected()
    current = PlanSettingsOut.model_validate(plan_settings.get_or_create(db, user))
    new_settings = _read_settings(book, current, out)
    categories = new_settings or current

    products = _read_products(book, out)
    transactions = _read_transactions(book, categories, out)
    fresh_products = _new_products(db, user, products)
    fresh_transactions = _new_transactions(db, user, transactions)

    if not dry_run:
        if new_settings is not None:
            stored = plan_settings.get_or_create(db, user)
            for key, value in new_settings.model_dump().items():
                setattr(stored, key, value)
        db.add_all(Product(user_id=user.id, **p.model_dump()) for p in fresh_products)
        db.add_all(Transaction(user_id=user.id, **t.model_dump()) for t in fresh_transactions)
        db.commit()

    return ImportReport(
        dry_run=dry_run,
        transactions_new=len(fresh_transactions),
        transactions_duplicated=len(transactions) - len(fresh_transactions),
        products_new=len(fresh_products),
        products_duplicated=len(products) - len(fresh_products),
        settings_updated=new_settings is not None,
        errors=out.errors,
        warnings=out.warnings,
    )
