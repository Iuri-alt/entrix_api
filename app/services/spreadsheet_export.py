"""Gera um .xlsx com o mesmo layout de entrada da planilha original (pode ser reimportado)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import PlanSettings, Product, Transaction, User
from app.services import plan_settings
from app.services import spreadsheet_layout as layout

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1B2A49")
_DATE_FORMAT = "DD/MM/YYYY"
_MONEY_FORMAT = '"R$" #,##0.00'


def _write_header(sheet: Worksheet, title: str, headers: list[str]) -> None:
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14)
    header_row = layout.FIRST_DATA_ROW - 1
    for col, text in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[cell.column_letter].width = max(14, len(text) + 4)


def _transactions_sheet(sheet: Worksheet, items: list[Transaction]) -> None:
    _write_header(sheet, "Lançamentos: receitas e despesas", layout.TRANSACTION_HEADERS)
    for row, t in enumerate(items, start=layout.FIRST_DATA_ROW):
        values = [
            t.date,
            t.description,
            layout.KIND_LABELS[t.kind],
            layout.GROUP_LABELS.get(t.expense_group or "", ""),
            t.category,
            float(t.amount),
            layout.STATUS_LABELS[t.status],
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row, column=col, value=value)
        sheet.cell(row=row, column=1).number_format = _DATE_FORMAT
        sheet.cell(row=row, column=6).number_format = _MONEY_FORMAT


def _products_sheet(sheet: Worksheet, items: list[Product]) -> None:
    _write_header(sheet, "Produtos que quero comprar", layout.PRODUCT_HEADERS)
    for row, p in enumerate(items, start=layout.FIRST_DATA_ROW):
        values = [
            p.name,
            p.category,
            float(p.price),
            layout.PRIORITY_LABELS[p.priority],
            p.target_date,
            float(p.saved),
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row, column=col, value=value)
        sheet.cell(row=row, column=3).number_format = _MONEY_FORMAT
        sheet.cell(row=row, column=5).number_format = _DATE_FORMAT
        sheet.cell(row=row, column=6).number_format = _MONEY_FORMAT


def _config_sheets(book: Workbook, settings: PlanSettings) -> None:
    config = book.create_sheet(layout.SHEET_CONFIG)
    config["A1"] = "Configurações do plano"
    config[layout.CONFIG_PLAN_START] = settings.plan_start
    config[layout.CONFIG_PLAN_END] = settings.plan_end
    config[layout.CONFIG_DAYS_SOON] = settings.days_soon
    config[layout.CONFIG_TIGHT_LIMIT] = settings.tight_limit
    for cell in (layout.CONFIG_PLAN_START, layout.CONFIG_PLAN_END):
        config[cell].number_format = _DATE_FORMAT
    for label, cell in zip(
        ["Primeiro mês de economia", "Prazo final do plano", "Dias para 'Próximo'", "Limite da meta apertada"],
        [layout.CONFIG_PLAN_START, layout.CONFIG_PLAN_END, layout.CONFIG_DAYS_SOON, layout.CONFIG_TIGHT_LIMIT],
    ):
        config[f"A{cell[1:]}"] = label

    first = layout.CONFIG_LIST_FIRST_ROW
    for offset, name in enumerate(settings.product_categories):
        config[f"{layout.CONFIG_PRODUCT_CATEGORIES_COL}{first + offset}"] = name
    for offset, name in enumerate(settings.income_categories + settings.expense_categories):
        config[f"{layout.CONFIG_TRANSACTION_CATEGORIES_COL}{first + offset}"] = name

    income = book.create_sheet(layout.SHEET_INCOME)
    income["A1"] = "Controle da renda"
    labels = ["Renda mensal", "Gastos fixos", "Gastos variáveis", "Outros compromissos"]
    for label, (key, cell) in zip(labels, layout.INCOME_BUDGET_CELLS.items()):
        income[f"A{cell[1:]}"] = label
        income[cell] = getattr(settings, key)
        income[cell].number_format = _MONEY_FORMAT
    income["A11"] = "Base usada no plano"
    income[layout.INCOME_PLAN_BASE] = layout.BASE_LABELS[settings.plan_base]

    simulation = book.create_sheet(layout.SHEET_SIMULATION)
    simulation["A5"] = "Economia por mês (R$)"
    for cell, value in zip(layout.SIMULATION_SCENARIO_CELLS, settings.scenarios):
        simulation[cell] = value


def export_workbook(db: Session, user: User) -> bytes:
    transactions = list(
        db.scalars(
            select(Transaction)
            .where(Transaction.user_id == user.id)
            .order_by(Transaction.date, Transaction.id)
        )
    )
    products = list(
        db.scalars(select(Product).where(Product.user_id == user.id).order_by(Product.id))
    )

    book = Workbook()
    book.active.title = layout.SHEET_PRODUCTS
    _products_sheet(book.active, products)
    _transactions_sheet(book.create_sheet(layout.SHEET_TRANSACTIONS), transactions)
    _config_sheets(book, plan_settings.get_or_create(db, user))

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
