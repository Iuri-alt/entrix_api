from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.services.spreadsheet_import import parse_money


def _workbook_bytes() -> bytes:
    book = Workbook()
    products = book.active
    products.title = "Produtos"
    products.append([])
    for _ in range(3):
        products.append([])
    products.append(["Produto", "Categoria", "Preço", "Prioridade", "Data desejada", "Já tenho guardado"])
    products.append(["Korg Kronos 2", "Música", 6000, "Média", datetime(2027, 1, 10), 110])
    products.append(["Ryzen 7 5200x", "Computador", 1020, "Alta", datetime(2026, 11, 10), 0])
    products.append(["Sem preço", "Outros", None, "Alta", None, 0])

    launches = book.create_sheet("Lançamentos")
    for _ in range(5):
        launches.append([])
    launches.append([datetime(2026, 10, 8), "Kung Fu", "Despesa", "Fixa", "Lazer", 160, "Pago"])
    launches.append([datetime(2026, 10, 8), "Aluguel garagem", "Despesa", "Fixa", "Outras receitas", 200, "Pendente"])
    launches.append([datetime(2026, 10, 7), "Salario", "Receita", "Variável", "Salário", 2568.19, "Pago"])
    launches.append([None, "Sem data", "Despesa", "", "Outros", 10, "Pago"])
    launches.append([datetime(2018, 5, 1), "Antigo", "Despesa", "", "Outros", 10, "Pago"])

    config = book.create_sheet("Config")
    config["B4"] = "=TODAY()"
    config["B5"] = datetime(2026, 10, 1)
    config["B6"] = datetime(2027, 12, 31)
    config["B7"] = 30
    config["B8"] = 0.7
    for i, name in enumerate(["Computador", "Música", "Outros"]):
        config[f"H{4 + i}"] = name
    for i, name in enumerate(["Salário", "Renda extra", "Outras receitas", "Lazer", "Outros"]):
        config[f"J{4 + i}"] = name

    renda = book.create_sheet("Renda")
    renda["B5"] = "R$ 2,568.19"
    renda["B6"] = 0
    renda["B11"] = "Previsto"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _upload(client, auth, dry_run: bool):
    files = {"file": ("planilha.xlsx", _workbook_bytes())}
    return client.post(f"/spreadsheet/import?dry_run={str(dry_run).lower()}", files=files, headers=auth)


@pytest.mark.parametrize(
    "raw, expected",
    [("R$ 2,568.19", 2568.19), ("R$ 2.568,19", 2568.19), ("120,5", 120.5), (75, 75.0), ("", None), ("abc", None)],
)
def test_parse_money(raw, expected):
    assert parse_money(raw) == expected


def test_dry_run_reports_without_writing(client, auth):
    report = _upload(client, auth, dry_run=True).json()
    assert report["transactions_new"] == 4
    assert report["products_new"] == 2
    assert [e["row"] for e in report["errors"]] == [8, 9]  # produto sem preço, lançamento sem data
    messages = " ".join(w["message"] for w in report["warnings"])
    assert "grupo ignorado" in messages
    assert "Outras receitas" in messages
    assert client.get("/transactions/", headers=auth).json() == []


def test_import_is_idempotent(client, auth):
    first = _upload(client, auth, dry_run=False).json()
    second = _upload(client, auth, dry_run=False).json()
    assert first["transactions_new"] == 4 and first["products_new"] == 2
    assert second["transactions_new"] == 0 and second["transactions_duplicated"] == 4
    assert second["products_new"] == 0 and second["products_duplicated"] == 2
    assert len(client.get("/transactions/", headers=auth).json()) == 4

    settings = client.get("/settings/", headers=auth).json()
    assert settings["budget_income"] == 2568.19
    assert settings["expense_categories"] == ["Lazer", "Outros"]
    income = next(t for t in client.get("/transactions/", headers=auth).json() if t["kind"] == "income")
    assert income["expense_group"] is None


def test_rejects_non_spreadsheet(client, auth):
    files = {"file": ("x.xlsx", b"nao sou planilha")}
    assert client.post("/spreadsheet/import", files=files, headers=auth).status_code == 400


def test_export_roundtrip(client, auth):
    _upload(client, auth, dry_run=False)
    response = client.get("/spreadsheet/export", headers=auth)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert {"Produtos", "Lançamentos", "Config", "Renda", "Simulação"} <= set(book.sheetnames)

    other = {"Authorization": auth["Authorization"]}
    files = {"file": ("export.xlsx", response.content)}
    report = client.post("/spreadsheet/import?dry_run=true", files=files, headers=other).json()
    assert report["transactions_new"] == 0 and report["transactions_duplicated"] == 4
    assert report["errors"] == []
